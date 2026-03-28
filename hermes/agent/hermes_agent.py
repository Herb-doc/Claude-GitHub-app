#!/usr/bin/env python3
"""
HERMES Agent — Health & Evidence Repository, Medical Expert & Synthesis
Google Drive Scanner + AI Parser for Daniel M. Phend, ND, MH

This script connects to Google Drive, scans medical documents,
extracts lab data, and maintains the master health record.
"""

import os
import io
import json
import sys
import tempfile
from datetime import datetime
from pathlib import Path

import anthropic
import fitz  # PyMuPDF
import pandas as pd
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from google.oauth2.credentials import Credentials
from google_auth_oauthlib.flow import InstalledAppFlow
from google.auth.transport.requests import Request
from googleapiclient.discovery import build
from googleapiclient.http import MediaIoBaseDownload

# ─── Configuration ──────────────────────────────────────────────────────────

SCOPES = [
    "https://www.googleapis.com/auth/drive.readonly",
    "https://www.googleapis.com/auth/documents.readonly",
    "https://www.googleapis.com/auth/drive.file",
]

SEARCH_KEYWORDS = [
    "blood", "lab", "result", "scan", "MRI", "CT", "ultrasound",
    "test", "wellness", "hormone", "testosterone", "immune",
    "ferritin", "hematocrit", "calprotectin", "diagnostic",
    "pathology", "radiology",
]

AGENT_DIR = Path(__file__).parent
PROJECT_DIR = AGENT_DIR.parent
DATA_FILE = PROJECT_DIR / "frontend" / "src" / "data" / "hermes_data.json"
CREDENTIALS_FILE = AGENT_DIR / "credentials.json"
TOKEN_FILE = AGENT_DIR / "token.json"

CLAUDE_MODEL = "claude-sonnet-4-20250514"

PARSE_SYSTEM_PROMPT = (
    "You are a medical data parser. Extract all lab values, "
    "reference ranges, dates, and flag any abnormals. Return "
    "structured JSON with fields: date, test_name, value, "
    "unit, reference_range, flag (CRITICAL/HIGH/LOW/WATCH/NORMAL), "
    "category, source. Return only valid JSON array, no commentary."
)


# ─── Step 1: Authentication ────────────────────────────────────────────────

def authenticate():
    """Authenticate with Google Drive using OAuth2."""
    creds = None

    if TOKEN_FILE.exists():
        creds = Credentials.from_authorized_user_file(str(TOKEN_FILE), SCOPES)

    if not creds or not creds.valid:
        if creds and creds.expired and creds.refresh_token:
            print("[HERMES] Refreshing expired token...")
            creds.refresh(Request())
        else:
            if not CREDENTIALS_FILE.exists():
                print(
                    "\n[HERMES] ERROR: credentials.json not found!\n"
                    f"  Expected location: {CREDENTIALS_FILE}\n\n"
                    "  To get this file:\n"
                    "  1. Go to https://console.cloud.google.com/\n"
                    "  2. Create a project (or select existing)\n"
                    "  3. Enable the Google Drive API and Google Docs API\n"
                    "  4. Go to Credentials > Create Credentials > OAuth Client ID\n"
                    "  5. Choose 'Desktop App' as the application type\n"
                    "  6. Download the JSON file and save it as credentials.json\n"
                    f"  7. Place it in: {AGENT_DIR}\n"
                )
                sys.exit(1)

            print("[HERMES] Opening browser for Google authentication...")
            flow = InstalledAppFlow.from_client_secrets_file(
                str(CREDENTIALS_FILE), SCOPES
            )
            creds = flow.run_local_server(port=0)

        with open(TOKEN_FILE, "w") as f:
            f.write(creds.to_json())
        print("[HERMES] Authentication successful. Token saved.")

    return creds


# ─── Step 2: Scan Drive ────────────────────────────────────────────────────

def scan_drive(service):
    """Search Google Drive for all medical documents using keywords."""
    print("\n[HERMES] Scanning Google Drive for medical documents...")
    all_files = {}

    for keyword in SEARCH_KEYWORDS:
        query = (
            f"fullText contains '{keyword}' and "
            f"(mimeType='application/pdf' or "
            f"mimeType='application/vnd.google-apps.document' or "
            f"mimeType='application/vnd.openxmlformats-officedocument.wordprocessingml.document') "
            f"and trashed=false"
        )

        try:
            results = service.files().list(
                q=query,
                pageSize=100,
                fields="files(id, name, mimeType, modifiedTime, webViewLink)",
                orderBy="modifiedTime desc",
            ).execute()

            for f in results.get("files", []):
                all_files[f["id"]] = f
        except Exception as e:
            print(f"  [WARN] Search for '{keyword}' failed: {e}")

    print(f"  Found {len(all_files)} unique medical documents.")
    return list(all_files.values())


# ─── Step 3: Extract Text ──────────────────────────────────────────────────

def extract_text_from_google_doc(service, file_id):
    """Extract text from a Google Doc."""
    try:
        doc = service.files().export(
            fileId=file_id, mimeType="text/plain"
        ).execute()
        return doc.decode("utf-8") if isinstance(doc, bytes) else doc
    except Exception as e:
        print(f"  [WARN] Could not read Google Doc {file_id}: {e}")
        return ""


def extract_text_from_pdf(service, file_id, file_name):
    """Download PDF and extract text with PyMuPDF."""
    try:
        request = service.files().get_media(fileId=file_id)
        buffer = io.BytesIO()
        downloader = MediaIoBaseDownload(buffer, request)

        done = False
        while not done:
            _, done = downloader.next_chunk()

        buffer.seek(0)
        with tempfile.NamedTemporaryFile(suffix=".pdf", delete=False) as tmp:
            tmp.write(buffer.read())
            tmp_path = tmp.name

        text = ""
        with fitz.open(tmp_path) as pdf:
            for page in pdf:
                text += page.get_text()

        os.unlink(tmp_path)
        return text
    except Exception as e:
        print(f"  [WARN] Could not extract PDF '{file_name}': {e}")
        return ""


def extract_text_from_docx(service, file_id, file_name):
    """Download Word doc and extract text with python-docx."""
    try:
        request = service.files().get_media(fileId=file_id)
        buffer = io.BytesIO()
        downloader = MediaIoBaseDownload(buffer, request)

        done = False
        while not done:
            _, done = downloader.next_chunk()

        buffer.seek(0)
        doc = Document(buffer)
        return "\n".join(p.text for p in doc.paragraphs)
    except Exception as e:
        print(f"  [WARN] Could not extract DOCX '{file_name}': {e}")
        return ""


def extract_documents(service, files):
    """Extract text from all found documents."""
    print("\n[HERMES] Extracting text from documents...")
    documents = []

    for f in files:
        file_id = f["id"]
        file_name = f["name"]
        mime_type = f["mimeType"]
        modified = f.get("modifiedTime", "")
        link = f.get("webViewLink", "")

        print(f"  Extracting: {file_name}")

        if mime_type == "application/vnd.google-apps.document":
            text = extract_text_from_google_doc(service, file_id)
        elif mime_type == "application/pdf":
            text = extract_text_from_pdf(service, file_id, file_name)
        elif "wordprocessingml" in mime_type:
            text = extract_text_from_docx(service, file_id, file_name)
        else:
            continue

        if text.strip():
            documents.append({
                "title": file_name,
                "date": modified,
                "content": text[:10000],  # Limit to prevent token overflow
                "source_url": link,
            })

    print(f"  Successfully extracted {len(documents)} documents.")
    return documents


# ─── Step 4: Parse with AI ─────────────────────────────────────────────────

def parse_with_ai(documents):
    """Send extracted text to Claude for structured parsing."""
    print("\n[HERMES] Parsing documents with Claude AI...")

    api_key = os.environ.get("ANTHROPIC_API_KEY")
    if not api_key:
        print(
            "\n[HERMES] ERROR: ANTHROPIC_API_KEY not set!\n"
            "  Set it with: export ANTHROPIC_API_KEY=your_key_here\n"
        )
        sys.exit(1)

    client = anthropic.Anthropic(api_key=api_key)
    all_findings = []

    for doc in documents:
        print(f"  Parsing: {doc['title']}")
        try:
            message = client.messages.create(
                model=CLAUDE_MODEL,
                max_tokens=4096,
                system=PARSE_SYSTEM_PROMPT,
                messages=[{
                    "role": "user",
                    "content": (
                        f"Document: {doc['title']}\n"
                        f"Date: {doc['date']}\n\n"
                        f"Content:\n{doc['content']}"
                    ),
                }],
            )

            response_text = message.content[0].text
            # Try to extract JSON from response
            try:
                findings = json.loads(response_text)
            except json.JSONDecodeError:
                # Try to find JSON array in response
                start = response_text.find("[")
                end = response_text.rfind("]") + 1
                if start >= 0 and end > start:
                    findings = json.loads(response_text[start:end])
                else:
                    print(f"    [WARN] Could not parse AI response for {doc['title']}")
                    continue

            # Add source info to each finding
            for finding in findings:
                finding["source"] = doc["title"]
                finding["source_url"] = doc.get("source_url", "")

            all_findings.extend(findings)
            print(f"    Found {len(findings)} lab values.")

        except Exception as e:
            print(f"    [WARN] AI parsing failed for {doc['title']}: {e}")

    print(f"\n  Total new findings parsed: {len(all_findings)}")
    return all_findings


# ─── Step 5: Deduplicate ───────────────────────────────────────────────────

def deduplicate(new_findings, existing_data):
    """Compare new findings against existing data. Skip exact duplicates."""
    existing_findings = existing_data.get("findings", [])

    # Build a set of existing keys for fast lookup
    existing_keys = set()
    for f in existing_findings:
        key = (f.get("date", ""), f.get("test_name", ""), str(f.get("value", "")))
        existing_keys.add(key)

    unique = []
    duplicates = 0

    for f in new_findings:
        key = (f.get("date", ""), f.get("test_name", ""), str(f.get("value", "")))
        if key in existing_keys:
            duplicates += 1
        else:
            unique.append(f)
            existing_keys.add(key)

    print(f"\n[HERMES] Deduplication: {len(unique)} new, {duplicates} duplicates skipped.")
    return unique


# ─── Step 6: Update Master Data ────────────────────────────────────────────

def update_master_data(new_findings):
    """Merge new findings into hermes_data.json."""
    print("\n[HERMES] Updating master health record...")

    if DATA_FILE.exists():
        with open(DATA_FILE) as f:
            existing_data = json.load(f)
    else:
        existing_data = {"findings": [], "protocols": [], "metadata": {}}

    unique_findings = deduplicate(new_findings, existing_data)

    existing_data["findings"].extend(unique_findings)
    existing_data["findings"].sort(key=lambda x: x.get("date", ""), reverse=True)
    existing_data["metadata"]["last_sync"] = datetime.now().isoformat()
    existing_data["metadata"]["total_findings"] = len(existing_data["findings"])

    with open(DATA_FILE, "w") as f:
        json.dump(existing_data, f, indent=2)

    print(f"  Master record updated: {len(existing_data['findings'])} total findings.")
    return existing_data, len(unique_findings)


# ─── Step 7: Generate Health Chart ─────────────────────────────────────────

def generate_health_chart(data):
    """Generate Daniel_Phend_Health_Chart.docx from master data."""
    print("\n[HERMES] Generating health chart document...")

    doc = Document()

    # Title
    title = doc.add_heading("Daniel M. Phend — Health Intelligence Chart", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph(
        f"Generated by HERMES on {datetime.now().strftime('%B %d, %Y at %I:%M %p')}"
    ).alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph("")

    # Group findings by category
    findings = data.get("findings", [])
    categories = {}
    for f in findings:
        cat = f.get("category", "Uncategorized")
        categories.setdefault(cat, []).append(f)

    for category in sorted(categories.keys()):
        doc.add_heading(category, level=1)

        table = doc.add_table(rows=1, cols=5)
        table.style = "Light Shading Accent 1"
        headers = table.rows[0].cells
        headers[0].text = "Date"
        headers[1].text = "Test"
        headers[2].text = "Value"
        headers[3].text = "Reference"
        headers[4].text = "Flag"

        for finding in sorted(categories[category], key=lambda x: x.get("date", "")):
            row = table.add_row().cells
            row[0].text = finding.get("date", "N/A")
            row[1].text = finding.get("test_name", "N/A")
            row[2].text = f"{finding.get('value', 'N/A')} {finding.get('unit', '')}"
            row[3].text = finding.get("reference_range", "N/A")
            row[4].text = finding.get("flag", "N/A")

        doc.add_paragraph("")

    # Protocols section
    protocols = data.get("protocols", [])
    if protocols:
        doc.add_heading("Active Protocols", level=1)
        for p in protocols:
            doc.add_paragraph(
                f"{p.get('name', 'Unknown')} — {p.get('purpose', '')}",
                style="List Bullet",
            )

    chart_path = PROJECT_DIR / "Daniel_Phend_Health_Chart.docx"
    doc.save(str(chart_path))
    print(f"  Health chart saved: {chart_path}")
    return chart_path


# ─── Step 8: Sync Report ───────────────────────────────────────────────────

def print_sync_report(docs_scanned, new_findings, duplicates_skipped, data):
    """Print a clean summary to terminal."""
    findings = data.get("findings", [])
    critical = [f for f in findings if f.get("flag") in ("CRITICAL", "URGENT")]

    print("\n" + "=" * 60)
    print("  HERMES SYNC REPORT")
    print("=" * 60)
    print(f"  Documents scanned:    {docs_scanned}")
    print(f"  New findings added:   {new_findings}")
    print(f"  Duplicates skipped:   {duplicates_skipped}")
    print(f"  Total findings:       {len(findings)}")

    if critical:
        print(f"\n  CRITICAL FLAGS ({len(critical)}):")
        for c in critical:
            print(f"    - {c.get('date', 'N/A')}: {c.get('test_name', '')} = "
                  f"{c.get('value', '')} {c.get('unit', '')} [{c.get('flag', '')}]")
    else:
        print("\n  No critical flags found.")

    print(f"\n  Health chart updated: Yes")
    print("=" * 60)
    print("\n  HERMES sync complete. Open dashboard to review new findings.")
    print("  Start dashboard: cd frontend && npm install && npm run dev\n")


# ─── Main ───────────────────────────────────────────────────────────────────

def main():
    print("\n" + "=" * 60)
    print("  HERMES — Personal Medical Intelligence")
    print("  Health & Evidence Repository, Medical Expert & Synthesis")
    print("=" * 60)

    # Step 1: Authenticate
    creds = authenticate()
    service = build("drive", "v3", credentials=creds)

    # Step 2: Scan Drive
    files = scan_drive(service)
    if not files:
        print("\n[HERMES] No medical documents found in Drive.")
        print("  Check that your Drive contains medical documents")
        print("  and that you granted read permissions.\n")
        return

    # Step 3: Extract text
    documents = extract_documents(service, files)
    if not documents:
        print("\n[HERMES] Could not extract text from any documents.\n")
        return

    # Step 4: Parse with AI
    all_findings = parse_with_ai(documents)

    # Step 5 & 6: Deduplicate and update
    data, new_count = update_master_data(all_findings)

    # Step 7: Generate health chart
    generate_health_chart(data)

    # Step 8: Report
    print_sync_report(
        docs_scanned=len(documents),
        new_findings=new_count,
        duplicates_skipped=len(all_findings) - new_count,
        data=data,
    )


if __name__ == "__main__":
    main()
